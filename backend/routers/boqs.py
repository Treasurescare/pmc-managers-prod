"""Wave 3 — BOQ (Bill of Quantities) + Contract Management."""
import io
from typing import List, Optional
from datetime import datetime, timezone
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from bson import ObjectId

from db import db, next_sequence
from models.boq import (
    BOQTemplateIn, BOQVendorSubmissionIn,
    ContractIn, VariationOrderIn,
)
from models.base import BaseDocument, now_iso
from security import require_role
from routers.vendor_auth import current_vendor
from audit import audit_log

router = APIRouter(prefix="/api/v1", tags=["boq-contracts"])


def _out(doc: dict) -> dict:
    return BaseDocument.out(doc)


# ---------- BOQ Templates ----------

@router.post("/boq/templates")
async def create_boq_template(payload: BOQTemplateIn,
                              admin=Depends(require_role("admin"))):
    doc = payload.model_dump()
    doc.update({"created_at": now_iso(), "created_by": admin["id"]})
    res = await db.boq_templates.insert_one(doc)
    doc["_id"] = res.inserted_id
    await audit_log("boq_template.create", admin, "boq_template", str(res.inserted_id),
                    {"name": payload.name, "items": len(payload.items)})
    return _out(doc)


@router.get("/boq/templates")
async def list_boq_templates(admin=Depends(require_role("admin", "reviewer", "auditor"))):
    # Bug #2 — return a shape frontend can normalise; also include tenant_id
    # in the future without breaking clients (currently BOQ templates are
    # global — no tenant_id yet — so filter is empty).
    items = [_out(t) async for t in db.boq_templates.find({}).sort("created_at", -1).limit(500)]
    return items


@router.get("/boq/templates/{template_id}")
async def get_boq_template(template_id: str, admin=Depends(require_role("admin", "reviewer", "auditor"))):
    doc = await db.boq_templates.find_one({"_id": ObjectId(template_id)})
    if not doc:
        raise HTTPException(status_code=404, detail="Not found")
    return _out(doc)


@router.post("/boq/templates/{template_id}/import-excel")
async def import_boq_excel(template_id: str, file: UploadFile = File(...),
                           admin=Depends(require_role("admin"))):
    """Import BOQ items from Excel — expects columns:
    Sl.No | Description | Unit | Quantity | Unit Rate | Remarks
    """
    from routers.boqs_helpers import rows_to_items
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    content = await file.read()
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="Empty spreadsheet")

    items, _skipped = rows_to_items(rows)

    r = await db.boq_templates.update_one(
        {"_id": ObjectId(template_id)},
        {"$set": {"items": items, "imported_at": now_iso()}},
    )
    if r.matched_count == 0:
        raise HTTPException(status_code=404, detail="Template not found")
    await audit_log("boq_template.import", admin, "boq_template", template_id,
                    {"items": len(items)})
    return {"ok": True, "items": len(items)}


# ---------- Vendor BOQ Submissions ----------

@router.post("/boq/submissions")
async def submit_boq(payload: BOQVendorSubmissionIn, vendor=Depends(current_vendor)):
    tmpl = await db.boq_templates.find_one({"_id": ObjectId(payload.template_id)})
    if not tmpl:
        raise HTTPException(status_code=404, detail="Template not found")
    bid = await db.bids.find_one({"_id": ObjectId(payload.bid_id), "vendor_id": vendor["id"]})
    if not bid:
        raise HTTPException(status_code=404, detail="Bid not found")

    # Compute total
    price_map = {int(p["sl_no"]): float(p.get("unit_rate", 0)) for p in payload.prices}
    total = 0.0
    priced_items = []
    for it in tmpl.get("items", []):
        rate = price_map.get(int(it.get("sl_no", 0)), 0.0)
        line_total = rate * float(it.get("quantity", 0))
        total += line_total
        priced_items.append({**it, "unit_rate": rate, "line_total": line_total})

    doc = {
        "template_id": payload.template_id,
        "bid_id": payload.bid_id,
        "vendor_id": vendor["id"],
        "notes": payload.notes,
        "priced_items": priced_items,
        "total_amount": round(total, 2),
        "currency": tmpl.get("currency", "USD"),
        "submitted_at": now_iso(),
    }
    res = await db.boq_submissions.insert_one(doc)
    doc["_id"] = res.inserted_id
    await audit_log("boq.submit", None, "bid", bid.get("bid_ref"),
                    {"total": doc["total_amount"]})
    return _out(doc)


@router.get("/boq/compare/{template_id}")
async def compare_boq_submissions(template_id: str,
                                   admin=Depends(require_role("admin", "committee_member"))):
    """Cross-vendor BOQ comparison for auto-comparison view."""
    tmpl = await db.boq_templates.find_one({"_id": ObjectId(template_id)})
    if not tmpl:
        raise HTTPException(status_code=404, detail="Not found")
    subs = [s async for s in db.boq_submissions.find({"template_id": template_id})]
    rows = []
    for item in tmpl.get("items", []):
        row = {"sl_no": item["sl_no"], "description": item["description"],
               "unit": item["unit"], "quantity": item["quantity"], "rates": {}}
        for s in subs:
            for pi in s.get("priced_items", []):
                if pi["sl_no"] == item["sl_no"]:
                    row["rates"][s["vendor_id"]] = pi.get("unit_rate", 0)
        rows.append(row)
    return {"template_id": template_id, "vendor_ids": [s["vendor_id"] for s in subs],
            "totals": [{"vendor_id": s["vendor_id"], "total": s["total_amount"]} for s in subs],
            "rows": rows}


# ---------- Contracts ----------

@router.post("/contracts")
async def create_contract(payload: ContractIn, admin=Depends(require_role("admin"))):
    seq = await next_sequence("contract_seq")
    year = datetime.now(timezone.utc).year
    contract_ref = payload.contract_ref or f"PMC-CON-{year}-{seq:04d}"

    # Verify the bid is actually awarded
    bid = await db.bids.find_one({"_id": ObjectId(payload.bid_id)})
    if not bid or bid.get("status") != "awarded":
        raise HTTPException(status_code=400, detail="Contract can only be created for an awarded bid")

    doc = payload.model_dump()
    doc.update({
        "contract_ref": contract_ref,
        "milestones": [m if isinstance(m, dict) else m.model_dump()
                       for m in payload.milestones],
        "variation_orders": [],
        "status": "active",
        "created_at": now_iso(),
        "created_by": admin["id"],
    })
    res = await db.contracts.insert_one(doc)
    doc["_id"] = res.inserted_id
    await audit_log("contract.create", admin, "contract", contract_ref,
                    {"value": payload.value, "bid": payload.bid_id})
    return _out(doc)


@router.get("/contracts")
async def list_contracts(admin=Depends(require_role("admin", "accounting", "reviewer", "auditor"))):
    # Bug #12 — return array; frontend already handles this.
    return [_out(c) async for c in db.contracts.find({}).sort("created_at", -1).limit(500)]


@router.get("/contracts/{contract_id}")
async def get_contract(contract_id: str,
                       admin=Depends(require_role("admin", "accounting", "reviewer", "auditor"))):
    doc = await db.contracts.find_one({"_id": ObjectId(contract_id)})
    if not doc:
        raise HTTPException(status_code=404, detail="Not found")
    return _out(doc)


@router.post("/contracts/{contract_id}/milestones/{order}/complete")
async def complete_milestone(contract_id: str, order: int,
                             admin=Depends(require_role("admin", "accounting"))):
    r = await db.contracts.update_one(
        {"_id": ObjectId(contract_id), "milestones.order": order},
        {"$set": {"milestones.$.status": "completed",
                  "milestones.$.completed_at": now_iso()}},
    )
    if r.matched_count == 0:
        raise HTTPException(status_code=404, detail="Milestone not found")
    await audit_log("milestone.complete", admin, "contract", contract_id, {"order": order})
    return {"ok": True}


@router.post("/contracts/variations")
async def add_variation(payload: VariationOrderIn, admin=Depends(require_role("admin"))):
    r = await db.contracts.update_one(
        {"_id": ObjectId(payload.contract_id)},
        {"$push": {"variation_orders": {**payload.model_dump(),
                                        "at": now_iso(), "by": admin["email"]}}},
    )
    if r.matched_count == 0:
        raise HTTPException(status_code=404, detail="Contract not found")
    await audit_log("contract.variation", admin, "contract", payload.contract_id,
                    {"delta": payload.amount_delta})
    return {"ok": True}
