"""Build annotated bug response Excel — mirrors the original bug list
(Sl No, Tab/Module, Bug/Observations, date, Posted By, Assigned To, Priority)
and appends two new columns: Status + Action Taken.
"""
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

SRC = Path("/tmp/originals/original_bugs.xlsx")
OUT = Path("/app/deliverables/PMC_Bugs_Response_v1.xlsx")
OUT.parent.mkdir(exist_ok=True)

# Response per bug_id (from Sl No column, expressed as int)
# Format: {id: (status, action_taken)}
RESPONSES = {
    1: ("FIXED — Live in prod",
        "Clone-as-draft now assigns 'order' field automatically in tender_templates.create_from_master. Shipped in backend-v5, live on www.pmcmanagers.com."),
    2: ("PARTIAL — In v11 zip · Retest",
        "Backend list_boq_templates() now capped at 500 with sort by created_at desc. If still not visible after v11/v15 deploy + hard refresh, likely a browser cache issue — please retest with Cmd/Ctrl+Shift+R."),
    3: ("FIXED — Live in prod",
        "Root cause: 403 responses from non-admin roles were breaking the Promise chain, forcing logout. Fixed with Promise.allSettled in BiddingOpsPanel.load. Shipped v5."),
    4: ("FIXED — Live in prod",
        "PDF generation now uses reportlab fallback when weasyprint fails. Blank form downloads working. Shipped v5."),
    5: ("FIXED — In v11 zip",
        "Root cause: reviewer role (madhavi) had create=False in DEFAULT_MATRIX. Fixed backend to allow reviewer.create=True but keep publish=False. Frontend now defaults is_published=false, disables publish checkbox for non-admins with clear label, defensive fallback if notice_ref is missing (no more 'undefined' toast). Ramya confirmed working."),
    6: ("OPEN — Awaiting repro",
        "Backend validation on mailing_address is standard. Please send screenshot showing exact field where error appears, so we can trace client-side validation."),
    7: ("OPEN — Wave E2",
        "Verify EMAIL_KEY env var propagates into Docker container. Add error log + toast on Resend failure. Scheduled next sprint."),
    8: ("CONFIRMED WORKING",
        "User (Kannan) confirmed 26-Aug: 'email is currently working, invite vendor is working properly'. No action required."),
    9: ("FIXED — Live in prod",
        "Defensive fallback added in AdminVendorCreate.submit — no more 'undefined' toast when response body is missing vendor_ref."),
    10: ("Duplicate of #1", "See Bug 1."),
    11: ("FIXED — Live in prod",
         "Added 45-second Promise.race timeout in L1ComparisonPanel.run. Page no longer hangs. Shipped v5."),
    12: ("PARTIAL — In v11 zip · Retest",
         "Backend list_contracts() capped at 500 with sort. If still not visible after v11 deploy + hard refresh, likely browser cache."),
    13: ("FIXED — Live in prod",
         "Promise.allSettled in MastersPanel.load — one failing sub-query no longer hangs the tab. Shipped v5."),
    14: ("FIXED — Live in prod",
         "Strict validation added in escrow.create_account: officer name min 2 chars, email required, roles must match side. Shipped v6."),
    15: ("OPEN — Wave E2",
         "Add default 'no anomalies detected' state + explain what the score checks. Scheduled next sprint."),
    16: ("FIXED — Live in prod",
         "NoticeIn.budget_amount now has ge=0 validator; frontend input has min='0' + NaN guard. Shipped v6."),
    17: ("OPEN — Wave E2",
         "DocuSign embedded-signing URL expiry — add retry page. Scheduled next sprint."),
    18: ("Merged with #16",
         "Negative budget rejection covered by Bug 16 fix."),
    19: ("OPEN — Retest first",
         "Backend and frontend look correct on code inspection. May be a caching bug — retest with hard refresh after v11/v15 deploy."),
    20: ("OPEN — Wave E2",
         "Show clear toast on invalid token at /register/status. Scheduled next sprint."),
    21: ("OPEN — Wave E2",
         "Need actual console error text from QA. Likely a benign CORS/401 preflight warning — please paste the exact error message."),
    22: ("FIXED — In v11 zip",
         "NoticesPanel.create success handler now performs full form reset instead of only clearing title/description."),
    23: ("FIXED — In v11 zip",
         "Toast now reads: 'Title is required (minimum 3 characters)' or 'Description is required (minimum 10 characters)' — explicit character count."),
    24: ("FIXED — In v11 zip",
         "Publish checkbox is now controlled by isMasterAdmin (me?.role === 'admin'). Master admin sees enabled checkbox with 'Publish immediately' label. Non-admins see disabled with 'reserved for Master Admin' label."),
    25: ("OPEN — Clarify UX",
         "Public corrigenda list already filters is_active=True (backend line 93). Admin list intentionally shows all with visual distinction. Please clarify with QA what visual state was expected — we may need stronger 'WITHDRAWN' badge."),
    26: ("OPEN — See #25",
         "Withdraw endpoint sets is_active=False + audit_log. Frontend refreshes list. Please retest and if still failing, capture browser Network tab response."),
    27: ("OPEN — Retest first",
         "registrations.decide backend endpoint updates status correctly with audit_log. Please retest with hard refresh."),
    28: ("FIXED — In v11 zip",
         "BulkInviteModal.submit now wrapped in try/catch with defensive result handling. Combined with root React ErrorBoundary — page no longer blanks or logs out. Modal stays open on error for retry."),
    29: ("OPEN — Awaiting repro",
         "Please capture screenshot showing what appears in the toast + Network tab response body."),
    30: ("OPEN — Retest first",
         "approve_and_create_vendor sets tenant_id correctly (registrations_helpers.py line 28). Backend list filters by admin.tenant_id. Should show — retest after hard refresh."),
    31: ("OPEN — Wave E2",
         "Audit i18n/*.json files — fill missing keys in ru/kz locales."),
    32: ("OPEN — Wave E2",
         "TopBar must refetch admin.me() on route change so avatar_data_url refreshes."),
    33: ("OPEN — Wave E2",
         "Add zod/yup validation to Section A/B/C: email regex + phone digits-only."),
    34: ("FIXED — In v11 zip",
         "Added cache-buster (?_=timestamp) query param on products list fetch — CDN/browser cache no longer serves stale data after save."),
    35: ("FIXED — In v11 zip",
         "Same fix as Bug 34 — cache-buster query param eliminates stale-data issue."),
    36: ("MITIGATED — In v11 zip",
         "Root React ErrorBoundary now catches audit-panel crashes — user stays logged in with soft error card. Full audit pagination fix scheduled for Wave E2."),
    37: ("Duplicate of #19",
         "See Bug 19."),
    38: ("OPEN — Wave E2",
         "Reset activeTemplate state on tab change (Tender/RFP/RFQ). Scheduled next sprint."),
    39: ("MITIGATED — In v11 zip",
         "Root React ErrorBoundary catches crash — session preserved. Full PDF-generation timeout guard scheduled for Wave E2."),
    40: ("OPEN — Wave E2",
         "L1 Comparison dropdown should filter for notices with awarded bids only. Scheduled next sprint."),
    41: ("OPEN — Wave E2",
         "Awards regenerate endpoint returns 200 but no visible change — add version stamp + confirmation toast. Scheduled next sprint."),
    42: ("OPEN — Retest first",
         "Backend + frontend look correct on inspection. Retest after v11/v15 deploy + hard refresh."),
    43: ("FIXED — In v11 zip",
         "Reverse auction start endpoint now returns explicit error: 'Cannot start auction — need at least 2 registered participants, currently N. Ask vendors to register first via the public link.'"),
    44: ("FIXED — In v11 zip",
         "Close endpoint now rejects close-on-draft with clear message. Idempotent if already closed."),
    45: ("OPEN — Retest first",
         "Backend admin_list filters by status parameter correctly. Please verify enum casing (must be 'rejected' lowercase) — retest."),
    46: ("FIXED — In v11 zip",
         "Nudge stale drafts now returns three-branch toast: 'No stale drafts found' / 'Checked N drafts — none met threshold' / 'Reminders sent to N of M'."),
    47: ("OPEN — Wave E2",
         "Add section-level validation gate in registration wizard — must fill required fields in Section A before advancing to B/C/D. Scheduled next sprint."),
    48: ("OPEN — Likely dup of #4",
         "Please verify with QA whether this differs from Bug 4 which was fixed. If a different flow, capture screenshot."),
    49: ("OPEN — Wave E2",
         "Stakeholder Reports snapshot returns non-JSON on error — wrap in try/catch, return JSON error body."),
    50: ("FIXED — In v11 zip",
         "Registration decide toast now reads 'Approval recorded' / 'Rejection recorded' / 'Clarification recorded' (fixed 'clarify recorded' typo)."),
    51: ("FIXED — In v11 zip",
         "send-reminders now inserts an admin notification per email sent (kind=registration.reminder_sent). Ops team now sees nudges in the notification centre."),
    52: ("Duplicate of #28",
         "See Bug 28. Fixed by ErrorBoundary + defensive submit."),
    53: ("", ""),  # No entry in original sheet at 53
    54: ("FIXED — In v11 zip",
         "Backend list_meetings now hides cancelled by default (include_cancelled param). Frontend has toggle. Cancelled meetings no longer clutter the panel."),
    55: ("Duplicate of #14",
         "See Bug 14. Escrow strict validation shipped in v6."),
    56: ("FIXED — In v11 zip",
         "ScoringRubricIn now validates: weights sum to 1.0 (±0.01), each max_score > 0, each weight ∈ [0,1], commercial_weight ∈ [0,1]. Returns 422 with clear message if invalid."),
    57: ("Duplicate of #39",
         "See Bug 39. Mitigated by ErrorBoundary; full fix scheduled Wave E2."),
    58: ("FIXED — In v11 zip",
         "Root React ErrorBoundary component wraps entire App — any panel crash now shows soft error card with 'Reload panel' + 'Try again' buttons. Session/auth is preserved."),
    59: ("", ""),  # 59 not in list
    60: ("Duplicate of #25",
         "See Bug 25. Awaiting UX clarification from QA."),
}

# Load original
wb = load_workbook(SRC)
ws = wb.active
last_col = ws.max_column
# Add two new columns
status_col = last_col + 1
action_col = last_col + 2

HEADER_FILL = PatternFill("solid", fgColor="0F1B3D")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
BORDER = Border(*[Side(style="thin", color="D0D0D0")] * 4)
FILL_FIXED = PatternFill("solid", fgColor="C6EFCE")
FILL_PARTIAL = PatternFill("solid", fgColor="FFEB9C")
FILL_OPEN = PatternFill("solid", fgColor="FFC7CE")
FILL_DUP = PatternFill("solid", fgColor="E7E6E6")

# Header row is row 2 (row 1 is title)
ws.cell(row=2, column=status_col, value="Status")
ws.cell(row=2, column=action_col, value="Action Taken / Response")
for c in (status_col, action_col):
    ws.cell(row=2, column=c).fill = HEADER_FILL
    ws.cell(row=2, column=c).font = HEADER_FONT
    ws.cell(row=2, column=c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(row=2, column=c).border = BORDER

# Fill responses
for r in range(3, ws.max_row + 1):
    sl_no = ws.cell(row=r, column=1).value
    try:
        bug_id = int(float(str(sl_no)))
    except Exception:
        bug_id = None
    if bug_id and bug_id in RESPONSES:
        status, action = RESPONSES[bug_id]
        ws.cell(row=r, column=status_col, value=status)
        ws.cell(row=r, column=action_col, value=action)
        # Colour by status keyword
        if "FIXED" in status or "CONFIRMED" in status:
            fill = FILL_FIXED
        elif "PARTIAL" in status or "MITIGATED" in status:
            fill = FILL_PARTIAL
        elif "Duplicate" in status:
            fill = FILL_DUP
        elif "OPEN" in status:
            fill = FILL_OPEN
        else:
            fill = None
        if fill:
            for c in (status_col, action_col):
                ws.cell(row=r, column=c).fill = fill
        for c in (status_col, action_col):
            ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=r, column=c).border = BORDER

# Column widths
ws.column_dimensions[get_column_letter(status_col)].width = 26
ws.column_dimensions[get_column_letter(action_col)].width = 85

wb.save(OUT)
print(f"Written {OUT}  size={OUT.stat().st_size / 1024:.1f} KB")
print(f"Original had {ws.max_row - 2} bug rows; responded to {sum(1 for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=status_col).value)} bugs.")
