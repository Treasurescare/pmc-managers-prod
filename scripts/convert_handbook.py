"""Convert the Trainee Handbook markdown into styled DOCX + PDF."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
import markdown as md_lib
from weasyprint import HTML, CSS

SRC = Path("/app/deliverables/PMC_Trainee_Handbook_v3.1.md")
OUT_DOCX = Path("/app/deliverables/PMC_Trainee_Handbook_v3.1.docx")
OUT_PDF = Path("/app/deliverables/PMC_Trainee_Handbook_v3.1.pdf")

md_text = SRC.read_text(encoding="utf-8")


# ---------- Inline **bold** parser ----------
def parse_inline(text):
    out = []
    idx = 0
    for m in re.finditer(r"\*\*(.+?)\*\*", text):
        if m.start() > idx:
            out.append((text[idx:m.start()], False))
        out.append((m.group(1), True))
        idx = m.end()
    if idx < len(text):
        out.append((text[idx:], False))
    return out or [(text, False)]


# ---------- DOCX ----------
def build_docx():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    NAVY = RGBColor(0x0F, 0x1B, 0x3D)
    GRAY = RGBColor(0x55, 0x55, 0x55)
    MUTED = RGBColor(0x77, 0x77, 0x77)

    def add_p(text, size=10, bold=False, italic=False, colour=None):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.size = Pt(size); r.font.name = "Calibri"
        if bold: r.font.bold = True
        if italic: r.font.italic = True
        if colour: r.font.color.rgb = colour

    def add_runs(p, text, size=10, bold_default=False):
        for chunk, is_bold in parse_inline(text):
            r = p.add_run(chunk)
            r.font.size = Pt(size); r.font.name = "Calibri"
            if is_bold or bold_default: r.font.bold = True

    # Cover
    add_p("PMC Managers v2", size=28, bold=True, colour=NAVY)
    add_p("Trainee Handbook · v3.1", size=16, colour=GRAY)
    add_p("Wave F revision · September 2026", size=10, italic=True, colour=MUTED)
    doc.add_paragraph(); doc.add_paragraph()

    for line in md_text.split("\n"):
        s = line.rstrip()
        if not s:
            continue
        if s.startswith("# ") and s[2:].lower().startswith("pmc managers"):
            continue  # already on cover
        if s.startswith("# "):
            add_p(s[2:], size=20, bold=True, colour=NAVY)
        elif s.startswith("## "):
            add_p(s[3:], size=15, bold=True, colour=NAVY)
        elif s.startswith("### "):
            add_p(s[4:], size=12, bold=True, colour=NAVY)
        elif s.startswith("---"):
            add_p("─" * 90, size=8, colour=MUTED)
        elif s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, s[2:])
        elif re.match(r"^\d+\.\s", s):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s+", "", s))
        else:
            p = doc.add_paragraph()
            add_runs(p, s)

    doc.save(OUT_DOCX)
    print(f"DOCX written: {OUT_DOCX}  ({OUT_DOCX.stat().st_size / 1024:.1f} KB)")


# ---------- PDF via markdown → HTML → weasyprint ----------
def build_pdf():
    html_body = md_lib.markdown(md_text, extensions=["tables", "fenced_code"])
    html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>PMC Trainee Handbook v3.1</title></head>
<body>
<div class="cover">
  <div class="brand">PMC Managers v2</div>
  <h1>Trainee Handbook · v3.1</h1>
  <div class="tag">Wave F revision · September 2026</div>
</div>
<div class="content">{html_body}</div>
</body></html>"""

    css = CSS(string="""
      @page { size: A4; margin: 22mm 20mm 20mm 20mm; }
      body { font-family: 'Helvetica','Arial',sans-serif; font-size: 10pt; color: #111; line-height: 1.5; }
      .cover { page-break-after: always; padding-top: 40mm; }
      .cover .brand { font-size: 22pt; font-weight: 700; color: #0F1B3D; letter-spacing: 1px; }
      .cover h1 { font-size: 32pt; font-weight: 800; color: #0F1B3D; margin: 6mm 0 2mm 0; }
      .cover .tag { font-size: 11pt; color: #666; font-style: italic; }
      h1 { color:#0F1B3D; font-size: 20pt; margin-top: 12mm; border-bottom: 1px solid #E7E6E6; padding-bottom: 3mm; }
      h2 { color:#0F1B3D; font-size: 14pt; margin-top: 8mm; }
      h3 { color:#0F1B3D; font-size: 11.5pt; margin-top: 5mm; }
      p, li { font-size: 10pt; }
      code { background: #F5F5F7; padding: 1px 4px; font-size: 9pt; border-radius: 2px; }
      hr { border: none; border-top: 1px solid #E7E6E6; margin: 8mm 0; }
      table { border-collapse: collapse; margin: 3mm 0; width: 100%; }
      th, td { border: 1px solid #D0D0D0; padding: 3mm 4mm; font-size: 9.5pt; text-align: left; }
      th { background: #0F1B3D; color: white; }
      strong { color: #0F1B3D; }
    """)
    HTML(string=html).write_pdf(str(OUT_PDF), stylesheets=[css])
    print(f"PDF written: {OUT_PDF}  ({OUT_PDF.stat().st_size / 1024:.1f} KB)")


if __name__ == "__main__":
    build_docx()
    build_pdf()
